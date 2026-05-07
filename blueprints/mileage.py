"""
Mileage blueprint — Phase 3 of the time-tracking spec.

Routes:

  GET   /command-deck/mileage/                  -- index page (filterable)
  GET   /command-deck/mileage/<id>/edit         -- edit form page
  GET   /command-deck/mileage/data              -- JSON for index page (filtered)
  POST  /command-deck/mileage/new               -- create
  POST  /command-deck/mileage/<id>/update       -- per-field patch
  POST  /command-deck/mileage/<id>/delete       -- hard delete
  POST  /command-deck/mileage/<id>/submit       -- toggle submitted_at
  POST  /command-deck/mileage/bulk-submit       -- mark many as submitted
  GET   /command-deck/mileage/export.xlsx       -- xlsx download (filtered)

Storage:
  - date: ISO YYYY-MM-DD (ET-local day)
  - submitted_at: ISO 8601 ET-local datetime, or NULL = pending reimbursement
  - rate_cents: snapshot of settings.reimbursement_rate_cents at entry time
  - miles: REAL, server stores whatever the form computes (round-trip math
    happens client-side; server doesn't recompute — trusts the client)
  - project_id: nullable in schema, required in v1 API (UI enforces)

Privacy:
  Default excludes is_private projects. Client opens with ?include_private=1
  after PIN unlock (same convention as reports).
"""
from datetime import date as _date, datetime
import io
import os
import pytz

from flask import (
	Blueprint, jsonify, redirect, render_template, request, send_file, url_for,
)

from helpers.auth import cd_auth_required, is_authenticated
from helpers.db import et_now, get_db


PRIVATE_PROJECTS_PIN = os.environ.get('PRIVATE_PROJECTS_PIN', '')


mileage_bp = Blueprint('mileage', __name__)


# ---- Helpers ----


def _parse_iso_date(s):
	"""Validate ISO YYYY-MM-DD; return the string if valid, else None."""
	if not s:
		return None
	try:
		_date.fromisoformat(s.strip())
		return s.strip()
	except (ValueError, TypeError, AttributeError):
		return None


def _to_int(v, default=None):
	if v in (None, '', 'null'):
		return default
	try:
		return int(v)
	except (ValueError, TypeError):
		return default


def _to_float(v, default=None):
	if v in (None, '', 'null'):
		return default
	try:
		return float(v)
	except (ValueError, TypeError):
		return default


def _to_bool(v):
	"""Form sends '1' / 'true' / 'on' for true; everything else false."""
	if v is None:
		return 0
	s = str(v).strip().lower()
	return 1 if s in ('1', 'true', 'on', 'yes', 'y') else 0


def _serialize_entry(row):
	d = dict(row)
	# Convenience derived fields for UI
	d['reimbursement_cents'] = int(round((d.get('miles') or 0) * (d.get('rate_cents') or 0)))
	d['is_submitted'] = bool(d.get('submitted_at'))
	return d


def _filters_from_request(args):
	"""Extract date/project/status/include-private filters from query string.
	Returns (clauses, params) for the WHERE clause + filter metadata."""
	clauses = []
	params = []

	start = _parse_iso_date(args.get('start'))
	end = _parse_iso_date(args.get('end'))
	if start:
		clauses.append('me.date >= ?')
		params.append(start)
	if end:
		clauses.append('me.date <= ?')
		params.append(end)

	project = args.get('project')
	if project:
		project_id = _to_int(project)
		if project_id:
			clauses.append('me.project_id = ?')
			params.append(project_id)

	status = args.get('status')
	if status == 'submitted':
		clauses.append('me.submitted_at IS NOT NULL')
	elif status == 'unsubmitted':
		clauses.append('me.submitted_at IS NULL')

	include_private = args.get('include_private') == '1'
	if not include_private:
		# Treat NULL project_id (future personal-trip rows) as non-private.
		clauses.append('(p.is_private IS NULL OR p.is_private = 0)')

	meta = {
		'start': start,
		'end': end,
		'project': project,
		'status': status or 'all',
		'include_private': include_private,
	}
	return clauses, params, meta


def _query_entries(conn, clauses, params):
	where_sql = ('WHERE ' + ' AND '.join(clauses)) if clauses else ''
	rows = conn.execute(f'''
		SELECT me.*,
		       p.title       AS project_title,
		       p.slug        AS project_slug,
		       p.is_private  AS project_is_private
		FROM mileage_entries me
		LEFT JOIN projects p ON me.project_id = p.id
		{where_sql}
		ORDER BY me.date DESC, me.id DESC
	''', params).fetchall()
	return rows


# ---- Index page ----


@mileage_bp.route('/command-deck/mileage/', methods=['GET'])
@mileage_bp.route('/command-deck/mileage', methods=['GET'])
@cd_auth_required
def mileage_index():
	conn = get_db()

	# Trackable projects for the form picker
	projects = conn.execute('''
		SELECT id, title, slug, is_private
		FROM projects
		WHERE project_type IN ('work_subproject', 'personal')
		  AND is_private = 0
		ORDER BY title ASC
	''').fetchall()

	settings_row = conn.execute('SELECT * FROM settings WHERE id = 1').fetchone()
	settings = dict(settings_row) if settings_row else {}

	conn.close()
	return render_template(
		'command_deck_mileage.html',
		projects=[dict(p) for p in projects],
		settings=settings,
		private_projects_enabled=bool(PRIVATE_PROJECTS_PIN),
	)


@mileage_bp.route('/command-deck/mileage/data', methods=['GET'])
def mileage_data():
	if not is_authenticated():
		return jsonify({'error': 'unauthorized'}), 403
	conn = get_db()
	clauses, params, meta = _filters_from_request(request.args)
	rows = _query_entries(conn, clauses, params)

	entries = [_serialize_entry(r) for r in rows]

	# Aggregate totals — overall, plus per-project and per-status splits.
	total_miles = sum((e.get('miles') or 0) for e in entries)
	total_cents = sum(e.get('reimbursement_cents', 0) for e in entries)
	unsubmitted_miles = sum((e.get('miles') or 0) for e in entries if not e.get('is_submitted'))
	unsubmitted_cents = sum(e.get('reimbursement_cents', 0) for e in entries if not e.get('is_submitted'))

	by_project = {}
	for e in entries:
		key = str(e.get('project_id') or 0)
		if key not in by_project:
			by_project[key] = {
				'title': e.get('project_title') or '(no project)',
				'miles': 0.0,
				'cents': 0,
				'count': 0,
			}
		by_project[key]['miles'] += (e.get('miles') or 0)
		by_project[key]['cents'] += e.get('reimbursement_cents', 0)
		by_project[key]['count'] += 1

	conn.close()
	return jsonify({
		'meta': meta,
		'entries': entries,
		'totals': {
			'count': len(entries),
			'miles': total_miles,
			'cents': total_cents,
			'unsubmitted_miles': unsubmitted_miles,
			'unsubmitted_cents': unsubmitted_cents,
			'by_project': by_project,
		},
	})


# ---- Edit page ----


@mileage_bp.route('/command-deck/mileage/<int:entry_id>/edit', methods=['GET'])
@cd_auth_required
def mileage_edit(entry_id):
	conn = get_db()
	row = conn.execute('''
		SELECT me.*, p.title AS project_title, p.slug AS project_slug
		FROM mileage_entries me
		LEFT JOIN projects p ON me.project_id = p.id
		WHERE me.id = ?
	''', (entry_id,)).fetchone()
	if not row:
		conn.close()
		return "Mileage entry not found", 404

	projects = conn.execute('''
		SELECT id, title, slug, is_private
		FROM projects
		WHERE project_type IN ('work_subproject', 'personal')
		  AND is_private = 0
		ORDER BY title ASC
	''').fetchall()
	settings_row = conn.execute('SELECT * FROM settings WHERE id = 1').fetchone()

	conn.close()
	return render_template(
		'command_deck_mileage_edit.html',
		entry=_serialize_entry(row),
		projects=[dict(p) for p in projects],
		settings=dict(settings_row) if settings_row else {},
	)


# ---- Create ----


@mileage_bp.route('/command-deck/mileage/new', methods=['POST'])
def mileage_new():
	if not is_authenticated():
		return jsonify({'error': 'unauthorized'}), 403

	data = request.form

	date_iso = _parse_iso_date(data.get('date'))
	if not date_iso:
		return jsonify({'error': 'invalid_date'}), 400

	miles = _to_float(data.get('miles'))
	if miles is None or miles < 0:
		return jsonify({'error': 'invalid_miles'}), 400

	conn = get_db()

	# Rate snapshot — caller can override (rare), otherwise use current setting
	override_rate = _to_int(data.get('rate_cents'))
	if override_rate is not None and override_rate >= 0:
		rate_cents = override_rate
	else:
		s = conn.execute('SELECT reimbursement_rate_cents FROM settings WHERE id = 1').fetchone()
		rate_cents = int(s['reimbursement_rate_cents']) if s else 67

	project_id = _to_int(data.get('project_id'))
	round_trip = _to_bool(data.get('round_trip'))
	odometer_start = _to_float(data.get('odometer_start'))
	odometer_end = _to_float(data.get('odometer_end'))
	vehicle = (data.get('vehicle') or 'a').strip().lower()
	if vehicle not in ('a', 'b'):
		vehicle = 'a'

	now = et_now()
	cur = conn.execute('''
		INSERT INTO mileage_entries
			(project_id, date, description, from_location, to_location,
			 round_trip, odometer_start, odometer_end, miles, rate_cents,
			 vehicle, notes, created, updated)
		VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
	''', (
		project_id,
		date_iso,
		(data.get('description') or '').strip() or None,
		(data.get('from_location') or '').strip() or None,
		(data.get('to_location') or '').strip() or None,
		round_trip,
		odometer_start,
		odometer_end,
		miles,
		rate_cents,
		vehicle,
		(data.get('notes') or '').strip() or None,
		now, now,
	))
	new_id = cur.lastrowid
	conn.commit()
	row = conn.execute('''
		SELECT me.*, p.title AS project_title, p.slug AS project_slug
		FROM mileage_entries me
		LEFT JOIN projects p ON me.project_id = p.id
		WHERE me.id = ?
	''', (new_id,)).fetchone()
	conn.close()

	# Form posts can ask for a redirect (HTML flow) or JSON (XHR flow). The
	# mileage form does both — XHR for in-page submit, HTML POST for the
	# /log-miles shortcut path.
	if data.get('return') == 'redirect':
		return redirect(url_for('mileage.mileage_index'))
	return jsonify({'success': True, 'entry': _serialize_entry(row)})


# ---- Update ----


@mileage_bp.route('/command-deck/mileage/<int:entry_id>/update', methods=['POST'])
def mileage_update(entry_id):
	if not is_authenticated():
		return jsonify({'error': 'unauthorized'}), 403
	data = request.get_json(silent=True) or request.form

	conn = get_db()
	row = conn.execute('SELECT * FROM mileage_entries WHERE id = ?', (entry_id,)).fetchone()
	if not row:
		conn.close()
		return jsonify({'error': 'not_found'}), 404

	updates = {}

	if 'date' in data:
		date_iso = _parse_iso_date(data.get('date'))
		if not date_iso:
			conn.close()
			return jsonify({'error': 'invalid_date'}), 400
		updates['date'] = date_iso

	if 'project_id' in data:
		val = data.get('project_id')
		updates['project_id'] = _to_int(val)  # None clears

	if 'description' in data:
		v = (data.get('description') or '').strip()
		updates['description'] = v or None

	if 'from_location' in data:
		v = (data.get('from_location') or '').strip()
		updates['from_location'] = v or None

	if 'to_location' in data:
		v = (data.get('to_location') or '').strip()
		updates['to_location'] = v or None

	if 'round_trip' in data:
		updates['round_trip'] = _to_bool(data.get('round_trip'))

	if 'odometer_start' in data:
		updates['odometer_start'] = _to_float(data.get('odometer_start'))

	if 'odometer_end' in data:
		updates['odometer_end'] = _to_float(data.get('odometer_end'))

	if 'miles' in data:
		miles = _to_float(data.get('miles'))
		if miles is None or miles < 0:
			conn.close()
			return jsonify({'error': 'invalid_miles'}), 400
		updates['miles'] = miles

	if 'rate_cents' in data:
		rate = _to_int(data.get('rate_cents'))
		if rate is None or rate < 0:
			conn.close()
			return jsonify({'error': 'invalid_rate'}), 400
		updates['rate_cents'] = rate

	if 'vehicle' in data:
		v = (data.get('vehicle') or '').strip().lower()
		if v not in ('a', 'b'):
			conn.close()
			return jsonify({'error': 'invalid_vehicle'}), 400
		updates['vehicle'] = v

	if 'notes' in data:
		v = (data.get('notes') or '').strip()
		updates['notes'] = v or None

	if not updates:
		conn.close()
		return jsonify({'error': 'no_fields'}), 400

	updates['updated'] = et_now()
	set_sql = ', '.join(f'{k} = ?' for k in updates)
	conn.execute(
		f'UPDATE mileage_entries SET {set_sql} WHERE id = ?',
		list(updates.values()) + [entry_id],
	)
	conn.commit()
	row = conn.execute('''
		SELECT me.*, p.title AS project_title, p.slug AS project_slug
		FROM mileage_entries me
		LEFT JOIN projects p ON me.project_id = p.id
		WHERE me.id = ?
	''', (entry_id,)).fetchone()
	conn.close()
	return jsonify({'success': True, 'entry': _serialize_entry(row)})


# ---- Delete ----


@mileage_bp.route('/command-deck/mileage/<int:entry_id>/delete', methods=['POST'])
def mileage_delete(entry_id):
	if not is_authenticated():
		return jsonify({'error': 'unauthorized'}), 403
	conn = get_db()
	row = conn.execute('SELECT id FROM mileage_entries WHERE id = ?', (entry_id,)).fetchone()
	if not row:
		conn.close()
		return jsonify({'error': 'not_found'}), 404
	conn.execute('DELETE FROM mileage_entries WHERE id = ?', (entry_id,))
	conn.commit()
	conn.close()
	return jsonify({'success': True, 'deleted_id': entry_id})


# ---- Submit toggle ----


@mileage_bp.route('/command-deck/mileage/<int:entry_id>/submit', methods=['POST'])
def mileage_submit_toggle(entry_id):
	if not is_authenticated():
		return jsonify({'error': 'unauthorized'}), 403

	data = request.get_json(silent=True) or request.form
	conn = get_db()
	row = conn.execute('SELECT * FROM mileage_entries WHERE id = ?', (entry_id,)).fetchone()
	if not row:
		conn.close()
		return jsonify({'error': 'not_found'}), 404

	# Allow explicit value (submitted=1/0) or implicit toggle.
	if 'submitted' in data:
		want = _to_bool(data.get('submitted'))
	else:
		want = 0 if row['submitted_at'] else 1

	new_submitted_at = et_now() if want else None
	conn.execute(
		'UPDATE mileage_entries SET submitted_at = ?, updated = ? WHERE id = ?',
		(new_submitted_at, et_now(), entry_id),
	)
	conn.commit()
	row = conn.execute('''
		SELECT me.*, p.title AS project_title, p.slug AS project_slug
		FROM mileage_entries me
		LEFT JOIN projects p ON me.project_id = p.id
		WHERE me.id = ?
	''', (entry_id,)).fetchone()
	conn.close()
	return jsonify({'success': True, 'entry': _serialize_entry(row)})


# ---- Bulk submit ----


@mileage_bp.route('/command-deck/mileage/bulk-submit', methods=['POST'])
def mileage_bulk_submit():
	if not is_authenticated():
		return jsonify({'error': 'unauthorized'}), 403
	data = request.get_json(silent=True) or {}
	ids = data.get('ids') or []
	if not isinstance(ids, list) or not ids:
		return jsonify({'error': 'ids_required'}), 400
	want = _to_bool(data.get('submitted', 1))

	int_ids = [int(x) for x in ids if str(x).isdigit()]
	if not int_ids:
		return jsonify({'error': 'ids_required'}), 400

	now = et_now()
	new_submitted_at = now if want else None

	conn = get_db()
	placeholders = ','.join('?' * len(int_ids))
	conn.execute(
		f'UPDATE mileage_entries '
		f'SET submitted_at = ?, updated = ? '
		f'WHERE id IN ({placeholders})',
		[new_submitted_at, now] + int_ids,
	)
	conn.commit()
	conn.close()
	return jsonify({'success': True, 'count': len(int_ids), 'submitted': bool(want)})


# ---- xlsx export ----


@mileage_bp.route('/command-deck/mileage/export.xlsx', methods=['GET'])
@cd_auth_required
def mileage_export_xlsx():
	# openpyxl is the only Phase-3 dependency added on top of the existing
	# stack. Imported lazily so the rest of the blueprint stays importable
	# on a server that hasn't installed it yet (logs a clear error instead
	# of failing at app startup).
	try:
		from openpyxl import Workbook
		from openpyxl.styles import Font, Alignment
	except ImportError:
		return (
			"openpyxl not installed on this server. "
			"On PythonAnywhere bash: pip3.10 install --user openpyxl",
			500,
		)

	conn = get_db()
	clauses, params, meta = _filters_from_request(request.args)
	rows = _query_entries(conn, clauses, params)
	conn.close()

	wb = Workbook()
	ws = wb.active
	ws.title = 'Mileage'

	headers = [
		'Date', 'Project', 'Description', 'From', 'To', 'Round Trip',
		'Odometer Start', 'Odometer End', 'Miles', 'Rate ($/mi)',
		'Reimbursement ($)', 'Vehicle', 'Notes', 'Submitted',
	]
	ws.append(headers)
	for cell in ws[1]:
		cell.font = Font(bold=True)
		cell.alignment = Alignment(horizontal='left')

	for r in rows:
		miles = float(r['miles'] or 0)
		rate_cents = int(r['rate_cents'] or 0)
		reimbursement = round(miles * rate_cents / 100.0, 2)
		ws.append([
			r['date'] or '',
			r['project_title'] or '',
			r['description'] or '',
			r['from_location'] or '',
			r['to_location'] or '',
			'Yes' if r['round_trip'] else 'No',
			r['odometer_start'] if r['odometer_start'] is not None else '',
			r['odometer_end'] if r['odometer_end'] is not None else '',
			miles,
			round(rate_cents / 100.0, 4),
			reimbursement,
			r['vehicle'] or '',
			r['notes'] or '',
			(r['submitted_at'][:10] if r['submitted_at'] else ''),
		])

	# Best-effort column widths so the file isn't ugly on first open.
	widths = [12, 24, 32, 22, 22, 10, 14, 14, 8, 10, 14, 8, 32, 12]
	for i, w in enumerate(widths, start=1):
		ws.column_dimensions[chr(64 + i) if i <= 26 else 'A' + chr(64 + i - 26)].width = w

	buf = io.BytesIO()
	wb.save(buf)
	buf.seek(0)

	# Filename reflects the active filter for "did I export the right month?"
	if meta['start'] and meta['end']:
		fname = f"mileage-{meta['start']}-to-{meta['end']}.xlsx"
	elif meta['start']:
		fname = f"mileage-from-{meta['start']}.xlsx"
	elif meta['end']:
		fname = f"mileage-to-{meta['end']}.xlsx"
	elif meta['status'] == 'unsubmitted':
		fname = 'mileage-unsubmitted.xlsx'
	else:
		fname = 'mileage-all.xlsx'

	return send_file(
		buf,
		mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
		as_attachment=True,
		download_name=fname,
	)
